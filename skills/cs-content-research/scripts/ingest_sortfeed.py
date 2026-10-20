#!/usr/bin/env python3
"""Turn raw SortFeed exports into internal reference records.

Usage:
  python ingest_sortfeed.py <exports_dir> <accounts.csv> <out reference_records.json> [--ratio 2]

exports_dir : raw SortFeed exports exactly as downloaded (.xlsx / .csv / .json),
              any mix of TikTok, Instagram Reels and Instagram Posts tabs.
accounts.csv: one row per account (see references/accounts_template.csv):
              platform,handle,role,followers,followers_source,competitor_source,approved
              role = client | competitor
              competitor_source = named_by_client | ai_found
              approved = yes | no | pending   (named_by_client counts as approved)

SortFeed exports have NO follower column, so followers come from accounts.csv.
A post passes when views >= ratio x followers (default 2). Missing views or
followers -> passes_ratio = null (never "viral" without data).

Prints a gate report and exits 2 if any ai_found competitor is not approved
(their rows are kept but marked BLOCKED_PENDING_APPROVAL).
"""
import csv, json, os, re, sys, datetime as dt

BAD = {'', '-', '—', 'n/a', 'na', '?', 'none', 'null'}


def num(v):
    """'14k' -> 14000, '4.5 M' -> 4500000, '1,234' -> 1234, 126749.0 -> 126749. Junk -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().lower().replace(' ', ' ')
    if s in BAD:
        return None
    s = re.sub(r'(?<=\d) (?=\d{3}\b)', '', s)           # 1 000 000
    m = re.fullmatch(r'(\d+),(\d+)\s*([mk])', s)          # 1,3m (EU decimal)
    if m:
        s = f'{m.group(1)}.{m.group(2)}{m.group(3)}'
    s = s.replace(',', '')
    m = re.fullmatch(r'(\d+(?:\.\d+)?)\s*([mk]?)\.?', s)
    if not m:
        return None
    x = float(m.group(1))
    return x * {'m': 1e6, 'k': 1e3}.get(m.group(2), 1)


def canon_link(u):
    u = (u or '').strip().split('?')[0].rstrip('/')
    u = re.sub(r'instagram\.com//+', 'instagram.com/', u)
    return u + '/' if 'instagram.com' in u else u


def handle(h):
    return (h or '').strip().lstrip('@').lower()


def read_table(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.json':
        data = json.load(open(path, encoding='utf8'))
        rows = data if isinstance(data, list) else data.get('data', [])
        return [dict(r) for r in rows]
    if ext == '.csv':
        with open(path, encoding='utf-8-sig', newline='') as f:
            return list(csv.DictReader(f))
    if ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        out = []
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            hdr = [str(h).strip() if h is not None else '' for h in rows[0]]
            out += [dict(zip(hdr, r)) for r in rows[1:] if any(c not in (None, '') for c in r)]
        return out
    return []


def detect(row):
    keys = set(row)
    if 'TikTok' in keys or 'code' in keys and 'shareCount' in keys:
        return 'tiktok', 'video'
    if 'Reel' in keys:
        return 'instagram', 'reel'
    if 'Post' in keys:
        return 'instagram', 'post'
    if 'userName' in keys and 'code' in keys:  # raw JSON objects
        return ('instagram', 'post' if row.get('mediaType') in (1, 8) else 'reel')
    return None, None


def g(row, *names):
    for n in names:
        if n in row and row[n] not in (None, ''):
            return row[n]
    return None


def to_record(row, src_file):
    platform, kind = detect(row)
    if not platform:
        return None
    user = handle(g(row, 'Profile', 'userName', 'Username'))
    link = g(row, 'TikTok', 'Reel', 'Post')
    if not link and row.get('code'):
        link = (f"https://www.tiktok.com/@{user}/video/{row['code']}" if platform == 'tiktok'
                else f"https://www.instagram.com/{user}/{'p' if kind == 'post' else 'reel'}/{row['code']}/")
    if kind == 'post' and link and '/reel/' in link:
        kind = 'reel'
    transcript = g(row, 'Transcript', 'transcript')
    return {
        'source_account': user,
        'platform': platform,
        'media_kind': kind,                      # video | reel | post (post = image or carousel; check visually)
        'reference_url': canon_link(link),
        'create_date': str(g(row, 'Create Date', 'Publish Date', 'createDate') or '') or None,
        'views': num(g(row, 'Views', 'viewCount')),
        'likes': num(g(row, 'Likes', 'likesCount')),
        'comments': num(g(row, 'Comments', 'commentsCount')),
        'shares': num(g(row, 'Shares', 'shareCount')),
        'saves': num(g(row, 'Saves', 'savesCount')),
        'outlier_score': num(g(row, 'Outlier Score', 'Outlier score')),
        'caption': g(row, 'Captions', 'Caption', 'caption'),
        'transcript': transcript,
        'transcript_status': 'SORTFEED' if transcript else 'TRANSCRIPT_MISSING',
        'source_provenance': f'SortFeed export: {os.path.basename(src_file)}',
    }


MIN_PASSING = 5


def main(exports_dir, accounts_csv, out_json, ratio=2.0, outlier=3.0):
    accounts = {}
    with open(accounts_csv, encoding='utf-8-sig', newline='') as f:
        for a in csv.DictReader(f):
            accounts[(a['platform'].strip().lower(), handle(a['handle']))] = a

    records, seen, skipped_files = [], set(), []
    for name in sorted(os.listdir(exports_dir)):
        p = os.path.join(exports_dir, name)
        if not os.path.isfile(p) or name.startswith(('.', '~$')):
            continue
        rows = read_table(p)
        got = 0
        for row in rows:
            r = to_record(row, p)
            if not r or not r['reference_url'] or r['reference_url'] in seen:
                continue
            seen.add(r['reference_url'])
            records.append(r)
            got += 1
        if not got:
            skipped_files.append(name)

    blocked, unknown_accounts = set(), set()
    for r in records:
        acc = accounts.get((r['platform'], r['source_account']))
        if acc is None:
            unknown_accounts.add((r['platform'], r['source_account']))
            r.update(client_or_competitor=None, followers=None, followers_source=None, approval_gate='ACCOUNT_NOT_IN_accounts.csv')
        else:
            r['client_or_competitor'] = acc['role'].strip().lower()
            r['followers'] = num(acc.get('followers'))
            r['followers_source'] = acc.get('followers_source') or None
            ok = (r['client_or_competitor'] == 'client' or acc.get('competitor_source', '').strip() == 'named_by_client'
                  or acc.get('approved', '').strip().lower() == 'yes')
            r['approval_gate'] = 'OK' if ok else 'BLOCKED_PENDING_APPROVAL'
            if not ok:
                blocked.add((r['platform'], r['source_account']))
        v, fo = r['views'], r['followers']
        r['views_to_followers'] = round(v / fo, 2) if v and fo else None
        r['passes_ratio'] = (v >= ratio * fo) if v is not None and fo else None
        r['viral_signal'] = 'RATIO' if r['passes_ratio'] else None
        r['eligibility_reason'] = (f'views {int(v):,} >= {ratio:g} x followers {int(fo):,}' if r['passes_ratio']
                                   else f'views {int(v):,} < {ratio:g} x followers {int(fo):,}' if r['passes_ratio'] is False
                                   else 'NO_DATA: views or followers missing')
        # filled later by the skill (visual check + transcript), never guessed here
        r.update(type=None, type_status='UNCLASSIFIED', type_evidence=None, hook=None, topic=None,
                 reaction_target_url=None, carousel_slides=None, asset_requirements=None,
                 approval_status='pending', used_status='No')

    # Exception, not a rule: only for an account with < 5 posts passing the ratio.
    # Candidates (views >= 1x followers, or >= outlier x account median) go to the CM; the 2x rule is never lowered otherwise.
    passing = {}
    for r in records:
        k = (r['platform'], r['source_account'])
        passing[k] = passing.get(k, 0) + (r['passes_ratio'] is True)
    for r in records:
        if r['passes_ratio'] is not False or passing[(r['platform'], r['source_account'])] >= MIN_PASSING:
            continue
        v, fo, o = r['views'], r['followers'], r.get('outlier_score')
        if v and fo and v >= fo:
            r['viral_signal'] = 'EXCEPTION_1X'
            r['eligibility_reason'] += f'; EXCEPTION: account has < {MIN_PASSING} posts at {ratio:g}x, views >= 1x followers (CM decision)'
        elif o is not None and o >= outlier:
            r['viral_signal'] = 'OUTLIER'
            r['eligibility_reason'] += f'; EXCEPTION: account has < {MIN_PASSING} posts at {ratio:g}x, views >= {outlier:g}x account median (CM decision)'

    records.sort(key=lambda r: (r['client_or_competitor'] or 'z', r['source_account'], -(r['views'] or 0)))
    meta = {'generated': dt.datetime.now().isoformat(timespec='seconds'), 'ratio_rule': f'views >= {ratio:g} x followers',
            'exception_rule': f'only if an account has < {MIN_PASSING} posts at {ratio:g}x: views >= 1x followers (EXCEPTION_1X) or >= {outlier:g}x account median (OUTLIER); CM decides',
            'exports_dir': os.path.abspath(exports_dir), 'records': len(records)}
    json.dump({'meta': meta, 'records': records}, open(out_json, 'w', encoding='utf8'), ensure_ascii=False, indent=1)

    by = {}
    for r in records:
        k = (r['platform'], r['source_account'])
        b = by.setdefault(k, [0, 0, 0, 0, 0])
        b[0] += 1; b[1] += r['passes_ratio'] is True; b[2] += r['passes_ratio'] is None; b[3] += r['transcript_status'] == 'SORTFEED'
        b[4] += r.get('viral_signal') in ('OUTLIER', 'EXCEPTION_1X')
    print(f'{len(records)} records -> {out_json}')
    print('account | posts | pass ratio | exception candidates | no data | transcripts')
    for (pl, h), b in sorted(by.items()):
        print(f'  {pl}:@{h} | {b[0]} | {b[1]} | {b[4]} | {b[2]} | {b[3]}')
        if b[1] < MIN_PASSING and b[4]:
            print(f'    EXCEPTION: @{h} has < {MIN_PASSING} posts passing {ratio:g}x followers; {b[4]} candidates (>= 1x followers or >= {outlier:g}x median) -> ask CM, then write_workbook --include-exceptions')
    for f in skipped_files:
        print('WARN: no SortFeed rows recognised in', f)
    for pl, h in sorted(unknown_accounts):
        print(f'WARN: @{h} ({pl}) is in the exports but not in accounts.csv -> add it with followers')
    if blocked:
        print('GATE: competitors found by AI need Content Manager approval before use:')
        for pl, h in sorted(blocked):
            print(f'  {pl}:@{h}')
        return 2
    return 0


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument('exports_dir'); ap.add_argument('accounts_csv'); ap.add_argument('out_json')
    ap.add_argument('--ratio', type=float, default=2.0)
    ap.add_argument('--outlier', type=float, default=3.0)
    a = ap.parse_args()
    sys.exit(main(a.exports_dir, a.accounts_csv, a.out_json, ratio=a.ratio, outlier=a.outlier))
